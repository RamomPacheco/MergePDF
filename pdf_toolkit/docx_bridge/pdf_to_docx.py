"""Convert PDF to DOCX and back for rich text editing."""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from pdf2docx import Converter


def _ensure_docx_path(path: Path) -> Path:
    """Return a path ending with .docx (required by docx2pdf)."""
    if path.suffix.lower() == ".docx":
        return path
    return path.with_suffix(".docx")


def _ensure_pdf_path(path: Path) -> Path:
    """Return a path ending with .pdf (required by docx2pdf)."""
    if path.suffix.lower() == ".pdf":
        return path
    return path.with_suffix(".pdf")


def _docx2pdf_convert(input_docx: Path, output_pdf: Path) -> None:
    """Convert DOCX to PDF via Word COM (docx2pdf)."""
    if sys.platform != "win32":
        raise RuntimeError("DOCX to PDF conversion is only supported on Windows.")

    try:
        import win32com.client  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(
            "pywin32 is required for Word conversion. "
            "Install it with: pip install pywin32"
        ) from exc

    from docx2pdf import convert as docx2pdf_convert

    src = _ensure_docx_path(input_docx.resolve())
    out = _ensure_pdf_path(output_pdf.resolve())
    out.parent.mkdir(parents=True, exist_ok=True)

    if not src.is_file():
        raise FileNotFoundError(f"File not found: {src}")

    try:
        docx2pdf_convert(str(src), str(out))
    except AssertionError as exc:
        raise RuntimeError(
            "Invalid paths for Word conversion. "
            "Use a .docx input and a .pdf output path."
        ) from exc
    except Exception as exc:
        message = str(exc).strip() or repr(exc)
        raise RuntimeError(
            f"Word failed to convert DOCX to PDF: {message}"
        ) from exc


def pdf_to_docx(input_file: str, output_file: str) -> None:
    """Convert a PDF file to DOCX format."""
    try:
        src = Path(input_file)
        if not src.is_file():
            raise FileNotFoundError(f"File not found: {src}")

        out = Path(output_file)
        out.parent.mkdir(parents=True, exist_ok=True)

        converter = Converter(str(src))
        try:
            converter.convert(str(out))
        finally:
            converter.close()
    except FileNotFoundError:
        raise
    except Exception as exc:
        raise RuntimeError(f"Failed to convert PDF to DOCX: {exc}") from exc


def docx_to_pdf(input_file: str, output_file: str) -> None:
    """Convert a DOCX file to PDF using Microsoft Word (docx2pdf)."""
    try:
        _docx2pdf_convert(Path(input_file), Path(output_file))
    except (FileNotFoundError, RuntimeError):
        raise
    except Exception as exc:
        raise RuntimeError(f"Failed to convert DOCX to PDF: {exc}") from exc


def _replace_in_paragraphs(paragraphs, old_text: str, new_text: str) -> None:
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)


def _replace_in_document(doc: Document, old_text: str, new_text: str) -> None:
    _replace_in_paragraphs(doc.paragraphs, old_text, new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_paragraphs(cell.paragraphs, old_text, new_text)


def edit_via_docx(
    input_pdf: str,
    output_pdf: str,
    old_text: str,
    new_text: str,
) -> None:
    """Edit PDF text by converting to DOCX, replacing, and converting back."""
    temp_dir: str | None = None
    try:
        if not old_text:
            raise ValueError("old_text cannot be empty")

        src = Path(input_pdf)
        if not src.is_file():
            raise FileNotFoundError(f"File not found: {src}")

        temp_dir = tempfile.mkdtemp()
        temp_docx = Path(temp_dir) / "edited.docx"

        pdf_to_docx(str(src), str(temp_docx))

        doc = Document(str(temp_docx))
        _replace_in_document(doc, old_text, new_text)
        doc.save(str(temp_docx))
        del doc

        out = _ensure_pdf_path(Path(output_pdf))
        out.parent.mkdir(parents=True, exist_ok=True)
        _docx2pdf_convert(temp_docx, out)
    except (ValueError, FileNotFoundError, RuntimeError):
        raise
    except Exception as exc:
        raise RuntimeError(f"Failed to edit PDF via DOCX: {exc}") from exc
    finally:
        if temp_dir and os.path.isdir(temp_dir):
            shutil.rmtree(temp_dir, ignore_errors=True)
